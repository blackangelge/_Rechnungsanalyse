"""Prüfregeln: Parser/Rechnen, Fehlwerte = nicht prüfbar, Einheiten-Umrechnung über Gruppen, Speichern an der Analyse, API."""

from copy import deepcopy
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient

from app.crud import schema_definition as sd
from app.database import get_db
from app.main import app
from app.models.schema_def import ExtractedRecord
from app.models.validation_rule import ValidationRule
from app.services import rule_engine as re_
from app.services.rule_engine import RuleError, RuleSpec
from app.services.rule_seed import seed_default_rules
from tests.test_schema_export import world  # noqa: F401
from tests.test_schema_extraction import SAMPLE


@pytest.fixture()
def client(sqlite_session_factory):
    def override():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _eval(expr, data=None, children=None):
    node = re_.Node(1, "beleg", data or {})
    for ent, kids in (children or {}).items():
        node.children[ent] = [re_.Node(10 + i, ent, d, i, node) for i, d in enumerate(kids)]
    return re_.Evaluator(node).eval(re_.parse(expr))


def test_arithmetic_precedence_and_functions():
    assert _eval("2 + 3 * 4") == 14
    assert _eval("(2 + 3) * 4") == 20
    assert _eval("-3 + 10 / 4") == Decimal("-0.5")
    assert _eval("ROUND(10 / 3; 2)") == Decimal("3.33")
    assert _eval("ABS(-1) + MAX(1; 7; 3) + MIN(4; 2)") == 10
    assert _eval("0.1 + 0.2") == Decimal("0.3")                                      # Decimal, kein Float-Rauschen


@pytest.mark.parametrize("bad", ["", "1 +", "foo(1)", "1,5 * 2", "2 $ 3", "(1", "COUNT(a.b)"])
def test_syntax_errors_are_reported(bad):
    with pytest.raises(RuleError):
        re_.parse(bad)


def test_no_eval_possible():
    with pytest.raises(RuleError):
        re_.parse("__import__('os').system('x')")


def test_missing_value_is_not_guessed_but_coalesce_helps():
    node = re_.Node(1, "position", {"menge": 5, "preis": None})
    with pytest.raises(re_.MissingValue) as exc:
        re_.Evaluator(node).eval(re_.parse("menge * preis"))
    assert exc.value.names == ["preis"]
    assert re_.Evaluator(node).eval(re_.parse("menge * COALESCE(preis; 2)")) == 10


def test_division_by_zero_makes_rule_unchecked():
    spec = RuleSpec(None, "x", "beleg", "1 / a", "1", )
    r = re_.evaluate_rule(spec, re_.Node(1, "beleg", {"a": 0}))
    assert r["status"] == "skipped" and "Division" in r["message"]


def test_sum_sumif_count():
    pos = [{"p": 10, "satz": 19}, {"p": 5, "satz": 7}, {"p": 20, "satz": 19}, {"p": None, "satz": 19}]
    kids = {"position": pos}
    assert _eval("SUM(position.p)", children=kids) == 35                              # Position ohne Wert wird ausgelassen
    assert _eval("SUMIF(position.p; position.satz; satz)", {"satz": 19}, kids) == 30
    assert _eval("SUMIF(position.p; position.satz; 7)", children=kids) == 5
    assert _eval("SUM(position.p * position.satz / 100)", children=kids) == Decimal("1.9") + Decimal("0.35") + Decimal("3.8")
    assert _eval("COUNT(position)", children=kids) == 4


def test_comparison_tolerances_and_comparators():
    node = re_.Node(1, "beleg", {"a": 100.0, "b": 100.04})
    assert re_.evaluate_rule(RuleSpec(None, "t", "beleg", "a", "b", tolerance_abs=0.05), node)["status"] == "ok"
    assert re_.evaluate_rule(RuleSpec(None, "t", "beleg", "a", "b", tolerance_abs=0.01), node)["status"] == "failed"
    assert re_.evaluate_rule(RuleSpec(None, "t", "beleg", "a", "b", tolerance_pct=0.1), node)["status"] == "ok"      # 0,1 % von 100
    assert re_.evaluate_rule(RuleSpec(None, "t", "beleg", "a", "b", comparator="<="), node)["status"] == "ok"
    assert re_.evaluate_rule(RuleSpec(None, "t", "beleg", "b", "a", comparator="<="), node)["status"] == "failed"


def test_definition_check_against_schema(world):
    v = world.version
    ok = RuleSpec(None, "t", "position", "menge * einzelpreis", "gesamtpreis")
    assert re_.check_rule_spec(ok, v) == []
    assert any("Feld „kaputt“" in p for p in re_.check_rule_spec(RuleSpec(None, "t", "position", "kaputt", "gesamtpreis"), v))
    assert any("Entität" in p for p in re_.check_rule_spec(RuleSpec(None, "t", "gibtsnicht", "1", "1"), v))
    assert any("kein Zahlenfeld" in p for p in re_.check_rule_spec(RuleSpec(None, "t", "position", "artikelbezeichnung", "1"), v))
    assert re_.check_rule_spec(RuleSpec(None, "t", "beleg", "SUM(lieferant.iban)", "1"), v)
    assert re_.check_rule_spec(RuleSpec(None, "t", "beleg", "SUM(position.gesamtpreis)", "gesamtbetrag_netto"), v) == []
    assert re_.check_rule_spec(RuleSpec(None, "t", "position", "eltern.gesamtbetrag_netto", "1"), v) == []
    assert re_.check_rule_spec(RuleSpec(None, "t", "beleg", "gesamtbetrag_netto", "1", comparator="!="), v)


def test_default_rules_on_analysed_document(world):
    seed_default_rules(world.db)
    world.db.commit()
    doc = world.add_doc()
    res = re_.evaluate_document(world.db, doc.id, store=False)
    by_name = {}
    for r in res["results"]:
        by_name.setdefault(r["name"], []).append(r)
    assert by_name["Summe der Positionen = Netto"][0]["status"] == "ok"               # 3.960 = 3.960
    assert by_name["Netto + Steuer = Brutto"][0]["status"] == "ok"
    pos = {r["position_index"]: r for r in by_name["Positionspreis = Menge × Einzelpreis"]}
    assert pos[0]["status"] == "ok"                                                   # 4400 × 90 / 100 = 3960
    assert pos[1]["status"] == "skipped" and "gesamtpreis" in pos[1]["message"]       # kein Positionspreis → nicht prüfbar
    assert res["summary"]["status"] == "ok"


def test_results_are_stored_during_analysis_and_wrong_totals_flag_review(world):
    seed_default_rules(world.db)
    world.db.commit()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"
    doc = world.add_doc(bad)
    root = world.db.query(ExtractedRecord).filter(ExtractedRecord.document_id == doc.id, ExtractedRecord.parent_record_id.is_(None)).one()
    assert root.data["_validation"]["status"] == "needs_review"
    failed = [c for c in root.data["_checks"] if c["status"] == "failed"]
    assert [c["name"] for c in failed] == ["Netto + Steuer = Brutto"] and failed[0]["diff"] == pytest.approx(-287.6)


def test_unit_conversion_through_groups(world):
    doc = world.add_doc()
    # Position 2: 25.880 kg → 25,88 t (Gruppe Masse, Basis t)
    root, _ = re_.load_tree(world.db, doc.id)
    kg_pos = root.children["position"][1]
    units = re_.load_unit_table(world.db)
    assert re_.Evaluator(kg_pos, units).eval(re_.parse("menge")) == Decimal("25880")
    assert re_.Evaluator(kg_pos, units).eval(re_.parse("menge.t")) == Decimal("25.88")
    assert re_.Evaluator(kg_pos, units).eval(re_.parse("menge.basis")) == Decimal("25.88")
    assert re_.Evaluator(kg_pos, units).eval(re_.parse("menge.g")) == Decimal("25880000")
    # Stück-Gruppe hat kein „t“ → nicht prüfbar statt falsch gerechnet
    piece = root.children["position"][0]
    with pytest.raises(re_.MissingValue):
        re_.Evaluator(piece, units).eval(re_.parse("menge.t"))
    spec = RuleSpec(None, "Tonnenpreis", "position", "menge.t * 412", "10662.56")
    assert re_.evaluate_rule(spec, kg_pos, units)["status"] == "ok"
    assert any("nicht angelegt" in p for p in re_.check_rule_spec(RuleSpec(None, "t", "position", "menge.furlong", "1"), world.version,
                                                                   re_.known_unit_keys(units)))


def test_evaluation_never_changes_stored_quantities(world):
    """Umrechnung gilt nur innerhalb der Formel; gespeicherte Mengen (Original + Einheit des LLM) bleiben unangetastet."""
    seed_default_rules(world.db)
    world.db.add(ValidationRule(name="t", entity_name="position", left_expr="menge.t", right_expr="menge.kg / 1000"))
    world.db.commit()
    doc = world.add_doc()

    def quantities():
        rows = world.db.query(ExtractedRecord).filter(ExtractedRecord.document_id == doc.id, ExtractedRecord.parent_record_id.isnot(None)).all()
        return {r.id: dict(r.data["menge"]) for r in rows if "menge" in r.data}
    before = quantities()
    re_.evaluate_document(world.db, doc.id)
    world.db.commit()
    world.db.expire_all()
    assert quantities() == before
    assert any(q["original_unit"] == "kg" and q["original_value"] == 25880 for q in before.values())


def test_rules_api_crud_test_and_run(client, world):
    doc = world.add_doc()
    assert client.post("/api/validation-rules/seed-defaults").json()["created"] == 4          # inkl. Varianten der Positionsregel
    assert client.post("/api/validation-rules/seed-defaults").json()["created"] == 0                      # nur wenn leer
    meta = client.get("/api/validation-rules/meta").json()
    pos = next(e for e in meta["entities"] if e["name"] == "position")
    assert {"menge", "einzelpreis", "gesamtpreis"} <= {f["name"] for f in pos["fields"]} and "kg" in meta["units"]
    beleg = next(e for e in meta["entities"] if e["name"] == "beleg")
    assert sorted(beleg["children"]) == ["position", "umsatzsteuer"] and all(e["name"] != "lieferant" for e in meta["entities"])

    body = {"name": "Eigene", "entity_name": "position", "left_expr": "menge.kg * 2", "right_expr": "einzelpreis", "severity": "warn"}
    bad = client.post("/api/validation-rules", json={**body, "left_expr": "nixda"})
    assert bad.status_code == 422 and "nixda" in bad.json()["detail"]
    rid = client.post("/api/validation-rules", json=body).json()["id"]
    t = client.post("/api/validation-rules/test", json={**body, "document_id": doc.id}).json()
    assert len(t["results"]) == 3 and {r["status"] for r in t["results"]} <= {"ok", "failed", "skipped"}
    assert client.put(f"/api/validation-rules/{rid}", json={**body, "enabled": False}).json()["enabled"] is False
    assert client.delete(f"/api/validation-rules/{rid}").status_code == 204

    assert client.get(f"/api/validation-rules/document/{doc.id}").json()["summary"]["status"] in ("ok", "warn")
    run = client.post("/api/validation-rules/run", json={}).json()
    assert run["documents"] == 1 and run["rules"] == 4 and sum(run["counts"].values()) == 1
    assert len(client.get("/api/validation-rules").json()) == 4


def test_document_list_shows_rule_status_and_filters(client, world):
    seed_default_rules(world.db)
    world.db.commit()
    good = world.add_doc()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"
    flagged = world.add_doc(bad)
    rows = {r["id"]: r for r in client.get("/api/documents").json()}
    assert rows[good.id]["rule_status"] == "ok" and rows[flagged.id]["rule_status"] == "needs_review"
    assert rows[flagged.id]["rule_failed"] == 1 and rows[good.id]["rule_skipped"] >= 1
    ids = client.get("/api/documents/ids", params={"validation_status": ["needs_review"]}).json()
    assert ids == [flagged.id]
    assert client.get("/api/documents/ids", params={"validation_status": ["ok"]}).json() == [good.id]


def _default_position_rule():
    from app.services.rule_seed import DEFAULT_RULES
    d = DEFAULT_RULES[0]
    return RuleSpec(None, d["name"], "position", d["left"], d["right"], tolerance_abs=d["tolerance_abs"], alternatives=d["alternatives"])


def _qty(db, value, unit):
    from app.services.units import UnitResolver, normalize_quantity
    return normalize_quantity(UnitResolver(db), {"value": value, "unit": unit})[0]


def test_default_position_rule_handles_price_per_tonne_and_per_vpe(world):
    """25.000 kg Schrott zu 420 € je Tonne; 4.230 Schrauben, 90 je VPE, 86 € je VPE."""
    units = re_.load_unit_table(world.db)
    rule = _default_position_rule()
    base = {"preis_je_menge": None, "preisnachlass_prozent": None, "preisnachlass_betrag": None}
    schrott = {**base, "menge": _qty(world.db, 25000, "kg"), "einzelpreis": 420, "gesamtpreis": 10500}
    r = re_.evaluate_rule(rule, re_.Node(1, "position", schrott), units)
    assert r["status"] == "ok" and r["variant"] == 2                                  # Variante „Basiseinheit“ (Tonne)
    wrong = {**schrott, "gesamtpreis": 25000 * 420}
    assert re_.evaluate_rule(rule, re_.Node(1, "position", wrong), units)["status"] == "ok" \
        and re_.evaluate_rule(rule, re_.Node(1, "position", wrong), units)["variant"] == 1      # Menge × Preis (je kg) wäre ebenfalls schlüssig
    assert re_.evaluate_rule(rule, re_.Node(1, "position", {**schrott, "gesamtpreis": 9999}), units)["status"] == "failed"

    schrauben = {**base, "menge": _qty(world.db, 4230, "Stück"), "verpackungsinhalt": _qty(world.db, 90, "Stück"),
                 "einzelpreis": 86, "gesamtpreis": 4042}
    r = re_.evaluate_rule(rule, re_.Node(1, "position", schrauben), units)
    assert r["status"] == "ok" and r["variant"] == 4                                  # Menge ÷ Inhalt je VPE × Preis
    # kein Verpackungsinhalt gelesen → 1: dann gilt Menge × Preis
    plain = {**base, "menge": _qty(world.db, 10, "Stück"), "einzelpreis": 5, "gesamtpreis": 50}
    r = re_.evaluate_rule(rule, re_.Node(1, "position", plain), units)
    assert r["status"] == "ok" and r["variant"] == 1
    bad = re_.evaluate_rule(rule, re_.Node(1, "position", {**schrauben, "gesamtpreis": 1000}), units)
    assert bad["status"] == "failed" and "Varianten" in bad["message"]


def test_variants_skipped_ones_do_not_block_and_all_missing_is_unchecked():
    spec = RuleSpec(None, "v", "beleg", "a * 2", "b", alternatives=["c * 2"])
    assert re_.evaluate_rule(spec, re_.Node(1, "beleg", {"a": 1, "b": 4, "c": 2}))["variant"] == 2     # Variante 1 falsch (2≠4), Variante 2 passt
    assert re_.evaluate_rule(spec, re_.Node(1, "beleg", {"a": 2, "b": 4}))["status"] == "ok"           # Variante 2 nicht auswertbar → egal
    r = re_.evaluate_rule(spec, re_.Node(1, "beleg", {"b": 4}))
    assert r["status"] == "skipped" and set(r["missing"]) == {"a", "c"}


def test_default_rule_is_upgraded_in_steps(world):
    """v1 (nur Menge × Preis) → streng mit Preiseinheit, sobald das Schema das Feld hat; Varianten-Fassung bleibt für Schemas ohne das Feld."""
    from app.services import rule_seed
    from app.services.rule_seed import DEFAULT_RULES, _OLD_POSITION_LEFT, seed_default_rules
    assert rule_seed.has_price_unit_field(world.version)
    world.db.add(ValidationRule(name=DEFAULT_RULES[0]["name"], entity_name="position", left_expr=_OLD_POSITION_LEFT, right_expr="gesamtpreis"))
    world.db.commit()
    assert seed_default_rules(world.db) == 1
    rule = world.db.query(ValidationRule).one()
    assert rule.left_expr == rule_seed._STRICT_LEFT and rule.alt_left_exprs == []
    assert seed_default_rules(world.db) == 0
    # vom Nutzer geänderte Regel bleibt unberührt
    rule.left_expr = "menge * einzelpreis"
    world.db.commit()
    assert seed_default_rules(world.db) == 0 and world.db.query(ValidationRule).one().left_expr == "menge * einzelpreis"
    # Schema ohne Feld → Varianten-Fassung
    field = next(f for f in next(e for e in world.version.entities if e.name == "position").fields if f.name == "preiseinheit")
    assert rule_seed.default_rules(world.version)[0]["alternatives"] == []
    world.db.delete(field)
    world.db.flush()
    world.db.expire(world.version)
    assert len(rule_seed.default_rules(world.version)[0]["alternatives"]) == 3


def test_workflow_node_check_rules_and_branch_on_rule_status(world):
    """Knoten „Regeln prüfen“ setzt den Prüfregel-Status; die Weiche kann darauf verzweigen."""
    import asyncio
    from app.services import workflow_runner as wr, workflow_service as ws
    seed_default_rules(world.db)
    world.db.commit()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"
    doc = world.add_doc(bad)
    from types import SimpleNamespace
    ctx = SimpleNamespace(doc_id=doc.id, schema_done=True, rule_status=None, node_io=None, document_type_id=1, document_type_name="Eingangsrechnung",
                          fields={}, known_supplier_name=None, page_count=1, invoice_page_count=1, validation=None)
    res = asyncio.run(wr._h_check_rules(ws.CONFIG_MODELS["check_rules"](), ctx))
    assert ctx.rule_status == "needs_review" and "weichen ab" in res.detail
    cfg = ws.CONFIG_MODELS["branch"](field="rule_status", op="equals", value="needs_review")
    assert asyncio.run(wr._h_branch(cfg, ctx)).handle == "true"
    ok_graph = {"nodes": [{"id": "s", "type": "start", "position": {"x": 0, "y": 0}, "data": {"label": "Start", "config": {}}},
                          {"id": "c", "type": "check_rules", "position": {"x": 0, "y": 0}, "data": {"label": "Regeln prüfen", "config": {}}},
                          {"id": "e", "type": "end", "position": {"x": 0, "y": 0}, "data": {"label": "Ende", "config": {}}}],
                 "edges": [{"id": "1", "source": "s", "target": "c", "sourceHandle": "out"}, {"id": "2", "source": "c", "target": "e", "sourceHandle": "out"}]}
    assert not ws.validate_graph(ok_graph).errors


def test_excel_export_can_include_rule_results(world):
    """Vorlage mit rule_check: Spalten „Prüfregel-Status“/„Prüfregeln“ (Beleg) und „Prüfregeln“ je Position, in beiden Layouts."""
    import io
    from openpyxl import load_workbook
    from app.models.export_template import ExportTemplate
    from app.services import schema_export
    seed_default_rules(world.db)
    world.db.commit()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"
    world.add_doc(bad)
    for layout in ("sheet_per_entity", "flat"):
        t = ExportTemplate(name=f"R-{layout}", layout=layout, quantity_mode="original",
                           config={"rule_check": True, "flat": {"root": "beleg", "child": "position"}})
        wb = load_workbook(io.BytesIO(schema_export.build_workbook(world.db, sd.get_active(world.db), t, schema_export.Filters())))
        if layout == "flat":
            rows = [[c.value for c in r] for r in wb.active.iter_rows()]
            head = rows[0]
            assert rows[1][head.index("Prüfregel-Status")] == "Prüfen!"
            assert "Netto + Steuer = Brutto" in rows[1][head.index("Prüfregeln (Beleg)")]
            assert any(h and h.startswith("Prüfregeln (Position)") for h in head)
        else:
            beleg = [[c.value for c in r] for r in wb["Beleg"].iter_rows()]
            assert beleg[1][beleg[0].index("Prüfregel-Status")] == "Prüfen!"
            pos = [[c.value for c in r] for r in wb["Position"].iter_rows()]
            col = pos[0].index("Prüfregeln")
            assert pos[1][col].startswith("OK") and "nicht prüfbar" in pos[2][col]       # Position 1 stimmt, Position 2 ohne Positionspreis
    plain = ExportTemplate(name="ohne", layout="flat", quantity_mode="original", config={"flat": {"root": "beleg", "child": "position"}})
    wb = load_workbook(io.BytesIO(schema_export.build_workbook(world.db, sd.get_active(world.db), plain, schema_export.Filters())))
    assert "Prüfregel-Status" not in [c.value for c in next(wb.active.iter_rows())]


def test_price_unit_conversion_in_formula(world):
    """menge.@preiseinheit: Menge in der gedruckten Preiseinheit (Tonne, VPE, leer = wie gedruckt)."""
    from app.services import rule_seed
    units = re_.load_unit_table(world.db)
    rule = RuleSpec(None, "p", "position", rule_seed._STRICT_LEFT, "gesamtpreis", tolerance_abs=0.02)
    base = {"preis_je_menge": None, "preisnachlass_prozent": None, "preisnachlass_betrag": None}

    def ev(menge, unit, preis, gesamt, pe, pack=None):
        data = {**base, "menge": _qty(world.db, *menge), "einzelpreis": preis, "gesamtpreis": gesamt, "preiseinheit": pe}
        if pack:
            data["verpackungsinhalt"] = _qty(world.db, *pack)
        return re_.evaluate_rule(rule, re_.Node(1, "position", data), units)

    # 1) 25.000 kg Schrott, 420 €/t → 25 t × 420 = 10.500; der Fehler „kg × Preis“ wird jetzt erkannt
    assert ev((25000, "kg"), None, 420, 10500, "t")["status"] == "ok"
    assert ev((25000, "kg"), None, 420, 10500000, "t")["status"] == "failed"
    # 2) 4.230 Schrauben, 90 je VPE, 86 € je VPE → 47 × 86 = 4.042
    assert ev((4230, "Stück"), None, 86, 4042, "VPE", pack=(90, "Stück"))["status"] == "ok"
    assert ev((4230, "Stück"), None, 86, 363780, "VPE", pack=(90, "Stück"))["status"] == "failed"
    assert ev((47, "Karton"), None, 86, 4042, "Karton")["status"] == "ok"                       # Menge schon in Verpackungen
    assert ev((10, "Sack"), None, 86, 860, "VPE")["status"] == "ok"                              # gleiche Gruppe, Faktor 1
    assert ev((4230, "Stück"), None, 86, 4042, "VPE")["status"] == "failed"                      # ohne Inhalt zählt 1 → 4230 × 86
    # leere Preiseinheit → Preis je gedruckter Einheit
    assert ev((10, "Stück"), None, 5, 50, None)["status"] == "ok"
    assert ev((25000, "kg"), None, 0.41, 10250, "")["status"] == "ok"
    # unbekannte oder unpassende Preiseinheit → nicht prüfbar statt geraten
    r = ev((10, "Stück"), None, 5, 50, "Furlong")
    assert r["status"] == "skipped" and "preiseinheit" in r["message"]
    assert ev((25000, "kg"), None, 420, 10500, "VPE")["status"] == "failed"                      # kg ≠ Verpackung: Inhalt 1 → 25.000 × 420
    # Definition: braucht das Textfeld im Schema
    assert re_.check_rule_spec(rule, world.version, re_.known_unit_keys(units)) == []
    assert any("Textfeld" in p for p in re_.check_rule_spec(RuleSpec(None, "x", "position", "menge.@nixda * einzelpreis", "gesamtpreis"), world.version))


def test_enable_price_unit_on_existing_schema(client, world):
    """Bestehendes Schema ohne Feld: Knopf legt Entwurf an, ergänzt das Feld, aktiviert und stellt die Standardregel um."""
    from app.services import schema_models
    from app.services.rule_seed import DEFAULT_RULES
    pos = next(e for e in world.version.entities if e.name == "position")
    world.db.delete(next(f for f in pos.fields if f.name == "preiseinheit"))
    world.db.commit()
    schema_models.clear_cache()
    assert seed_default_rules(world.db) == 4
    world.db.commit()
    assert client.get("/api/validation-rules/meta").json()["price_unit_field"] is False
    assert len(client.get("/api/validation-rules").json()[0]["alt_left_exprs"]) == 3

    r = client.post("/api/validation-rules/enable-price-unit").json()
    assert r["changed"] is True and r["rules_updated"] == 1 and r["schema_version"] == 2
    assert client.get("/api/validation-rules/meta").json()["price_unit_field"] is True
    first = client.get("/api/validation-rules").json()[0]
    assert first["alt_left_exprs"] == [] and "@preiseinheit" in first["left_expr"] and first["problems"] == []
    assert client.post("/api/validation-rules/enable-price-unit").json()["changed"] is False          # idempotent


def test_filter_rule_issues_mismatch_and_not_calculable(world):
    from app.crud.document import get_ids_filtered
    seed_default_rules(world.db)
    world.db.commit()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"                     # Summe stimmt nicht
    d_bad = world.add_doc(bad)
    d_skip = world.add_doc()                                    # Position ohne Positionspreis → nicht prüfbar
    clean = deepcopy(SAMPLE)
    clean["positionen"] = clean["positionen"][:1]
    d_ok = world.add_doc(clean)
    world.db.commit()
    mismatch = get_ids_filtered(world.db, rule_issues=["mismatch"])
    notcalc = get_ids_filtered(world.db, rule_issues=["notcalc"])
    both = get_ids_filtered(world.db, rule_issues=["mismatch", "notcalc"])
    assert d_bad.id in mismatch and d_skip.id not in mismatch and d_ok.id not in mismatch
    assert d_skip.id in notcalc and d_ok.id not in notcalc
    assert {d_bad.id, d_skip.id} <= set(both) and d_ok.id not in both


def test_background_run_with_progress_cancel_and_parallel(world, sqlite_session_factory):
    from app.services import rule_run
    seed_default_rules(world.db)
    world.db.commit()
    bad = deepcopy(SAMPLE)
    bad["gesamtbetrag_brutto"] = "5.000,00"
    ids = [world.add_doc(bad).id, world.add_doc().id, world.add_doc().id]
    world.db.commit()
    # Ergebnisse veralten lassen, dann neu prüfen
    for rec in world.db.query(ExtractedRecord).filter(ExtractedRecord.parent_record_id.is_(None)).all():
        d = dict(rec.data); d.pop("_validation", None); d.pop("_checks", None); rec.data = d
    world.db.commit()
    for workers in (1, 3):
        rule_run._state.clear(); rule_run._state.update(running=True, total=0, processed=0, counts={}, flagged=[], errors=[], error_count=0)
        st = rule_run.run(sqlite_session_factory, None, workers)
        assert st["running"] is False and st["total"] == 3 and st["processed"] == 3 and st["error_count"] == 0
        assert st["counts"].get("needs_review") == 1 and ids[0] in st["flagged"]
    world.db.expire_all()
    root = world.db.query(ExtractedRecord).filter(ExtractedRecord.document_id == ids[0], ExtractedRecord.parent_record_id.is_(None)).one()
    assert root.data["_validation"]["status"] == "needs_review"
    # Abbruch vor dem Start der Blöcke: nichts mehr verarbeitet
    rule_run._state.clear(); rule_run._state.update(running=True, total=0, processed=0, counts={}, flagged=[], errors=[], error_count=0)
    rule_run._cancel.set()
    orig = rule_run._cancel.clear
    rule_run._cancel.clear = lambda: None
    try:
        st = rule_run.run(sqlite_session_factory, None, 1)
    finally:
        rule_run._cancel.clear = orig
        rule_run._cancel.clear()
    assert st["processed"] == 0 and st["cancelled"] is True


def test_price_basis_examples_from_the_prompt_rules(world):
    """Beispiel 1: 4.400 St, 90 € je VPE, 100 Schrauben = 1 VPE → Menge ÷ Preis-je-Menge × Einzelpreis; Beispiel 2: 25.500 kg, 420 €/Tonne."""
    from app.services import rule_seed
    units = re_.load_unit_table(world.db)
    rule = RuleSpec(None, "p", "position", rule_seed._STRICT_LEFT, "gesamtpreis", tolerance_abs=0.02)

    def ev(menge, preis, gesamt, pe, per=None, pack=None):
        data = {"preis_je_menge": per, "preisnachlass_prozent": None, "preisnachlass_betrag": None, "menge": _qty(world.db, *menge),
                "einzelpreis": preis, "gesamtpreis": gesamt, "preiseinheit": pe}
        if pack:
            data["verpackungsinhalt"] = _qty(world.db, *pack)
        return re_.evaluate_rule(rule, re_.Node(1, "position", data), units)

    assert ev((4400, "St"), 90, 3960, "VPE", per=100)["status"] == "ok"                       # (4400 / 100) × 90
    assert ev((4400, "St"), 90, 396000, "VPE", per=100)["status"] == "failed"
    assert ev((4400, "St"), 90, 3960, "VPE", per=100, pack=(100, "St"))["status"] == "ok"     # Zahl doppelt angegeben → nicht doppelt geteilt
    assert ev((4400, "St"), 90, 3960, "VPE", pack=(100, "St"))["status"] == "ok"              # nur als Inhalt je VPE angegeben
    assert ev((25500, "kg"), 420, 10710, "Tonne")["status"] == "ok"                           # 25,5 t × 420
    assert ev((25500, "kg"), 420, 10710000, "Tonne")["status"] == "failed"
