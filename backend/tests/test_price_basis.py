"""Preisbasis (Preis je N Einheiten), Verpackungsinhalt und Einheitentyp: KI-Auslesung, Prüfung, Nachrüstung, Artikelsuche, Excel."""

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import crud
from app.database import get_db
from app.main import app
from app.models.import_batch import ImportBatch
from app.models.order_position import OrderPosition
from app.services import ai_service, plausibility_service as ps, validation_backfill as vb


@pytest.fixture()
def client(sqlite_session_factory, monkeypatch):
    for target in ("app.services.change_log.SessionLocal", "app.services.validation_backfill.SessionLocal"):
        monkeypatch.setattr(target, sqlite_session_factory)

    def _get_db():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _get_db
    yield TestClient(app)
    app.dependency_overrides.clear()


def _doc(db, positions, raw=None):
    b = ImportBatch(company_name="F", year=2025, import_folder_path="/x", storage_folder_path="/s", status="done")
    db.add(b)
    db.commit()
    d = crud.document.create(db, b.id, "a.pdf", 1)
    crud.document.save_extraction(db, d.id, {"supplier_name": "L", "invoice_number": "R-1"}, positions, json.dumps(raw or {}))
    d.status = "done"
    db.commit()
    return d


def _schema_doc(db, positions):
    """Beleg mit Positionen im NEUEN Schema (extracted_records) — für die Artikelsuche."""
    from app.crud import schema_definition as sd
    from app.models.schema_def import ExtractedRecord
    from app.models.units import UnitGroup
    from app.services.schema_seed import seed_initial_schema
    from app.services.units import UnitResolver, normalize_quantity
    from app.services.units_seed import seed_default_units

    if db.query(UnitGroup).first() is None:
        seed_default_units(db)
    if sd.get_active(db) is None:
        seed_initial_schema(db)
    version = sd.get_active(db)
    beleg_e = next(e for e in version.entities if e.name == "beleg")
    pos_e = next(e for e in version.entities if e.name == "position")
    resolver = UnitResolver(db)
    b = ImportBatch(company_name="F", year=2025, import_folder_path="/x", storage_folder_path="/s", status="done")
    db.add(b)
    db.commit()
    d = crud.document.create(db, b.id, "a.pdf", 1)
    crud.document.update_status(db, d.id, "done")
    root = ExtractedRecord(document_id=d.id, entity_id=beleg_e.id, schema_version_id=version.id, parent_record_id=None,
                           data={"rechnungsnummer": "R-1", "_master": {"lieferant": {"name": "L"}}, "_validation": {"status": "ok"}})
    db.add(root)
    db.flush()
    for i, p in enumerate(positions):
        menge, _ = normalize_quantity(resolver, {"value": p.get("quantity"), "unit": p.get("unit")})
        pack = None
        if p.get("pack_size") is not None:
            pack = {"original_value": p["pack_size"], "original_unit": p.get("pack_unit"),
                    "value": None, "unit": None, "status": "ok"}
        db.add(ExtractedRecord(
            document_id=d.id, entity_id=pos_e.id, schema_version_id=version.id, parent_record_id=root.id, position_index=i,
            data={"artikelbezeichnung": p.get("product_description"), "artikelnummer": p.get("article_number"),
                  "menge": menge, "einzelpreis": p.get("unit_price"), "gesamtpreis": p.get("total_price"),
                  "preis_je_menge": p.get("price_quantity"), "verpackungsinhalt": pack,
                  "mengeneinheit_typ": p.get("unit_type")},
        ))
    db.commit()
    return d


# ── KI-Auslesung ────────────────────────────────────────────────────────────

def test_ai_mapping_reads_basis_pack_and_type():
    data = {"positionen": [
        {"artikelbezeichnung": "Schraube", "menge": 4400, "mengeneinheit": "Stück", "einzelpreis": 90, "gesamtpreis": 3960, "preis_je_menge": 100,
         "mengeneinheit_typ": "Stück"},
        {"artikelbezeichnung": "Mutter", "menge": 44, "mengeneinheit": "VPE", "einzelpreis": 90, "gesamtpreis": 3960,
         "verpackungsinhalt": 100, "verpackungsinhalt_einheit": "Stück", "mengeneinheit_typ": "VPE"},
        {"artikelbezeichnung": "Dübel", "menge": 500, "mengeneinheit": "St", "einzelpreis": 12, "gesamtpreis": 60, "preiseinheit": "je 100"},
        {"artikelbezeichnung": "Stahl", "menge": 2, "mengeneinheit": "t", "einzelpreis": 400, "gesamtpreis": 800, "preiseinheit": "t"},
    ]}
    _, pos = ai_service._map_new_format(data)
    assert (pos[0]["price_quantity"], pos[0]["unit_type"]) == (100.0, "stueck")
    assert (pos[1]["pack_size"], pos[1]["pack_unit"], pos[1]["unit_type"], pos[1]["price_quantity"]) == (100.0, "Stück", "verpackung", None)
    assert pos[2]["price_quantity"] == 100.0           # aus der Preiseinheit „je 100“
    assert pos[3]["price_quantity"] is None            # „t“ ist keine Zählangabe


def test_prompt_rules_are_appended_once_and_mention_json_keys():
    out = ai_service.with_unit_rules("Basisprompt")
    for key in ("preis_je_menge", "verpackungsinhalt", "verpackungsinhalt_einheit", "mengeneinheit_typ", "NEBENKOSTEN"):
        assert key in out
    for key in ("KEINE zusätzliche Position", "GENAU SO", "nur, wenn sie gedruckt sind"):
        assert key in out
    assert ai_service.with_unit_rules(out) == out


# ── Plausibilitätsprüfung ───────────────────────────────────────────────────

def _validate(pos, raw_pos=None):
    return ps.validate_invoice({}, [pos], {"positionen": [raw_pos or {}]})


def test_price_per_100_is_plausible_with_stored_basis_or_ai_hint():
    base = {"quantity": 4400, "unit": "Stück", "unit_price": 90, "total_price": 3960}
    assert not [i for i in _validate({**base, "price_quantity": 100}).issues if i.code == "position_mismatch"]
    assert not [i for i in _validate(dict(base), {"preis_je_menge": 100}).issues if i.code == "position_mismatch"]
    # ohne Basis: Abweichung wird weiter gemeldet (bisheriges Verhalten)
    assert any(i.code in ("position_mismatch", "price_unit_differs") for i in _validate(dict(base)).issues)


def test_position_calc_explains_basis():
    res = _validate({"quantity": 4400, "unit": "Stück", "unit_price": 90, "total_price": 3960, "price_quantity": 100})
    calcs = res.to_dict()["position_calcs"]
    assert any("Preisbasis" in step for step in calcs[0]["steps"]) and calcs[0]["status"] == "ok"


# ── Nachrüstung ohne KI ─────────────────────────────────────────────────────

def test_backfill_detects_basis_from_totals_and_pack_from_text(client, db):
    d = _doc(db, [
        {"product_description": "Schraube M4", "quantity": 4400, "unit": "Stück", "unit_price": 90, "total_price": 3960},
        {"product_description": "Dübel (VE = 250 Stück)", "quantity": 4, "unit": "VPE", "unit_price": 30, "total_price": 120},
        {"product_description": "Stahl", "quantity": 2, "unit": "t", "unit_price": 400, "total_price": 800},
        {"product_description": "Rätsel", "quantity": 3, "unit": "Stück", "unit_price": 10, "total_price": 31},     # nicht erklärbar
    ])
    assert vb.count_pending(price_basis=True) == 1
    vb.start(price_basis=True, background=False)
    db.expire_all()
    p = {x.position_index: x for x in db.query(OrderPosition).all()}
    assert float(p[0].price_quantity) == 100 and p[0].unit_type == "stueck"
    assert float(p[1].pack_size) == 250 and p[1].pack_unit == "Stück" and p[1].unit_type == "verpackung" and p[1].price_quantity is None
    assert p[2].unit_type == "gewicht" and p[2].price_quantity is None
    assert p[3].price_quantity is None                                                   # nichts geraten
    # zweiter Lauf ändert nichts
    vb.start(price_basis=True, background=False)
    db.expire_all()
    assert float(db.query(OrderPosition).filter_by(position_index=0).one().price_quantity) == 100


# ── Artikelsuche und Excel ──────────────────────────────────────────────────

def test_article_search_returns_content_values(client, db):
    _schema_doc(db, [{"product_description": "Schraube", "quantity": 44, "unit": "VPE", "unit_price": 90, "total_price": 3960,
                      "pack_size": 100, "pack_unit": "Stück", "unit_type": "verpackung"},
                     {"product_description": "Mutter", "quantity": 4400, "unit": "Stück", "unit_price": 90, "total_price": 3960, "price_quantity": 100},
                     {"product_description": "Ohne", "quantity": 3, "unit": "Stück", "unit_price": 2, "total_price": 6}])
    rows = {r["product_description"]: r for r in client.get("/api/positions").json()}
    s = rows["Schraube"]
    assert float(s["content_quantity"]) == 4400 and s["content_unit"] == "Stück" and round(float(s["content_unit_price"]), 4) == 0.9
    assert float(s["quantity"]) == 44 and s["unit"] == "VPE"                  # Originalwerte unverändert
    m = rows["Mutter"]
    assert float(m["content_quantity"]) == 4400 and round(float(m["content_unit_price"]), 4) == 0.9 and float(m["price_quantity"]) == 100
    assert rows["Ohne"]["content_quantity"] is None


def test_article_search_pack_in_weight_units_gets_base_values(client, db):
    _schema_doc(db, [{"product_description": "Sack", "quantity": 40, "unit": "Sack", "unit_price": 10, "total_price": 400, "pack_size": 25, "pack_unit": "kg"}])
    r = client.get("/api/positions").json()[0]
    assert r["base_unit"] == "t" and round(float(r["base_quantity"]), 4) == 1.0 and round(float(r["base_unit_price"]), 2) == 400.0


def test_excel_contains_new_columns_with_stored_values(client, db):
    d = _doc(db, [{"product_description": "Schraube", "quantity": 4400, "unit": "Stück", "unit_price": 90, "total_price": 3960,
                   "price_quantity": 100, "pack_size": 100, "pack_unit": "Stück", "unit_type": "stueck"}])
    batch_id = d.batch_id
    r = client.get(f"/api/imports/{batch_id}/export")
    ws = load_workbook(io.BytesIO(r.content))["Positionen"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    col = dict(zip(header, row))
    assert col["Preis je (Einheiten)"] == 100 and col["Inhalt je Verpackung"] == 100
    assert col["Einheit des Inhalts"] == "Stück" and col["Einheitentyp"] == "stueck"
    assert col["Gesamtpreis (€)"] == 3960 and col["Menge"] == 4400                      # unverändert, keine Berechnung
