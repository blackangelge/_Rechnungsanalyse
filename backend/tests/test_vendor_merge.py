"""Lieferanten (vendor) zusammenführen + Dublettensuche (API, SQLite)."""

import json

import pytest
from fastapi.testclient import TestClient

from app import crud
from app.database import get_db
from app.main import app
from app.models.import_batch import ImportBatch
from app.models.vendor import Vendor
from app.models.vendor_bank_account import VendorBankAccount
from app.models.workflow import Workflow


@pytest.fixture()
def client(sqlite_session_factory):
    def _get_db():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _get_db
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture()
def world(db):
    """Drei Lieferanten (zwei Dubletten + ein fremder) mit Belegen; Belege hängen am Namen."""
    batch = ImportBatch(company_name="T", year=2026, import_folder_path="/x", storage_folder_path="/s/x")
    db.add(batch)
    db.commit()
    a = Vendor(name="ACME GmbH", city="Berlin")
    b = Vendor(name="Acme GmbH & Co. KG", street="Weg 1", postal_code="10115", vat_id="DE123456789")
    c = Vendor(name="Fremd AG")
    a.bank_accounts = [VendorBankAccount(iban="DE02120300000000202051", bank_name="Bank A")]
    b.bank_accounts = [VendorBankAccount(iban="DE02 1203 0000 0000 2020 51", bank_name="Bank B"),
                       VendorBankAccount(iban="DE89370400440532013000")]
    db.add_all([a, b, c])
    db.commit()

    def doc(name):
        d = crud.document.create(db, batch.id, f"{name}{len(db.new)}.pdf", 1)
        crud.document.save_extraction(db, d.id, {"supplier_name": name}, [], "{}")
        return d
    docs = {"a1": doc("ACME GmbH"), "a2": doc("ACME GmbH"), "b1": doc("Acme GmbH & Co. KG"), "c1": doc("Fremd AG")}

    class W:
        pass
    w = W()
    w.db, w.a, w.b, w.c, w.docs = db, a, b, c, docs
    return w


def _vendor_ids(db):
    db.expire_all()
    return sorted(v.id for v in db.query(Vendor).all())


def _name(w, key):
    w.db.expire_all()
    return w.docs[key].extraction.vendor_id


def test_duplicate_groups_endpoint(client, world):
    groups = client.get("/api/vendors/duplicate-groups").json()
    assert len(groups) == 1
    g = groups[0]
    assert [s["id"] for s in g["suppliers"]] == [world.a.id, world.b.id]
    assert g["confidence"] == "high" and g["reasons"][0] == "iban"          # gleiche IBAN trotz Leerzeichen
    assert "gleiche IBAN" in g["reason_labels"] and g["suppliers"][0]["document_count"] == 2


def test_duplicate_groups_use_every_bank_account(client, world):
    d = Vendor(name="Völlig anders KG")
    d.bank_accounts = [VendorBankAccount(iban="DE89 3704 0044 0532 0130 00")]       # = zweite IBAN von b
    world.db.add(d)
    world.db.commit()
    g = client.get("/api/vendors/duplicate-groups", params={"name_check": "false"}).json()
    assert len(g) == 1 and [s["id"] for s in g[0]["suppliers"]] == [world.a.id, world.b.id, d.id]
    # mit Namensprüfung (Standard) bleibt der namentlich fremde Lieferant draußen
    g = client.get("/api/vendors/duplicate-groups").json()
    assert [s["id"] for s in g[0]["suppliers"]] == [world.a.id, world.b.id]


def test_duplicate_groups_criteria_filter(client, world):
    assert client.get("/api/vendors/duplicate-groups", params={"criteria": ["vat_id"]}).json() == []
    g = client.get("/api/vendors/duplicate-groups", params={"criteria": ["iban"]}).json()
    assert len(g) == 1 and g[0]["reasons"] == ["iban"]
    assert client.get("/api/vendors/duplicate-groups", params={"criteria": ["bogus"]}).status_code == 422


def test_merge_fills_gaps_moves_bank_accounts_and_documents(client, world):
    r = client.post("/api/vendors/merge", json={"target_id": world.a.id, "source_ids": [world.b.id]})
    assert r.status_code == 200, r.text
    body = r.json()
    s = body["supplier"]
    assert body["merged_ids"] == [world.b.id] and body["moved_documents"] == 1 and body["moved_bank_accounts"] == 1
    assert s["id"] == world.a.id and s["document_count"] == 3
    assert s["name"] == "ACME GmbH" and s["city"] == "Berlin"                # Ziel behält seine Werte
    assert s["street"] == "Weg 1" and s["vat_id"] == "DE123456789"          # Leeres wird gefüllt
    assert sorted(ba["iban"] for ba in s["bank_accounts"]) == ["DE02120300000000202051", "DE89370400440532013000"]
    assert _vendor_ids(world.db) == sorted([world.a.id, world.c.id])
    assert _name(world, "b1") == "ACME GmbH"                                 # Beleg gehört jetzt zum Ziel
    assert _name(world, "c1") == "Fremd AG"                                  # fremde Belege unberührt
    assert world.db.query(VendorBankAccount).count() == 2                    # Duplikat-IBAN entfällt, nichts verwaist
    assert client.get("/api/vendors/duplicate-groups").json() == []


def test_explicit_fields_win_and_target_documents_follow_new_name(client, world):
    r = client.post("/api/vendors/merge", json={
        "target_id": world.a.id, "source_ids": [world.b.id],
        "fields": {"name": "  Acme GmbH & Co. KG ", "city": "Hamburg"},
    }).json()
    assert r["supplier"]["name"] == "Acme GmbH & Co. KG" and r["supplier"]["city"] == "Hamburg"
    assert r["moved_documents"] == 2                                         # die beiden ACME-Belege; b1 hatte den Namen schon
    assert {_name(world, k) for k in ("a1", "a2", "b1")} == {"Acme GmbH & Co. KG"}
    assert r["supplier"]["document_count"] == 3


def test_documents_keep_names_without_flag(client, world):
    r = client.post("/api/vendors/merge", json={
        "target_id": world.a.id, "source_ids": [world.b.id], "update_document_names": False}).json()
    assert r["moved_documents"] == 0 and _name(world, "b1") == "Acme GmbH & Co. KG"


def test_merge_multiple_sources_and_workflow_carry_over(client, world):
    wf = Workflow(name="WF", graph=json.dumps({"nodes": [], "edges": []}))
    world.db.add(wf)
    world.db.commit()
    world.b.workflow_id = wf.id
    world.db.commit()
    r = client.post("/api/vendors/merge", json={"target_id": world.a.id, "source_ids": [world.c.id, world.b.id, world.b.id]})
    assert r.status_code == 200
    assert r.json()["merged_ids"] == [world.c.id, world.b.id]                # Duplikate in der Liste ignoriert
    assert r.json()["supplier"]["workflow_id"] == wf.id and r.json()["moved_documents"] == 2
    assert _vendor_ids(world.db) == [world.a.id]


def test_merge_keeps_target_workflow(client, world):
    w1, w2 = (Workflow(name=n, graph=json.dumps({"nodes": [], "edges": []})) for n in ("W1", "W2"))
    world.db.add_all([w1, w2])
    world.db.commit()
    world.a.workflow_id, world.b.workflow_id = w1.id, w2.id
    world.db.commit()
    r = client.post("/api/vendors/merge", json={"target_id": world.a.id, "source_ids": [world.b.id]}).json()
    assert r["supplier"]["workflow_id"] == w1.id


@pytest.mark.parametrize("payload, status, text", [
    ({"target_id": 1, "source_ids": []}, 422, None),
    ({"target_id": 1, "source_ids": [1]}, 400, "nicht gleichzeitig"),
    ({"target_id": 9999, "source_ids": [1]}, 404, "Ziel-Lieferant #9999"),
    ({"target_id": 1, "source_ids": [8888]}, 404, "#8888"),
    ({"target_id": 1, "source_ids": [2], "fields": {"unbekannt": "x"}}, 400, "Unbekannte Felder"),
    ({"target_id": 1, "source_ids": [2], "fields": {"name": " "}}, 400, "Name darf nicht leer"),
])
def test_merge_validation_errors_change_nothing(client, world, payload, status, text):
    before = _vendor_ids(world.db)
    r = client.post("/api/vendors/merge", json=payload)
    assert r.status_code == status
    if text:
        assert text in r.json()["detail"]
    assert _vendor_ids(world.db) == before


def test_merge_is_atomic_on_failure(world, monkeypatch):
    real_delete = world.db.delete

    def boom(obj):
        raise RuntimeError("db down")
    monkeypatch.setattr(world.db, "delete", boom)
    with pytest.raises(RuntimeError):
        crud.vendor.merge(world.db, world.a.id, [world.b.id])
    monkeypatch.setattr(world.db, "delete", real_delete)
    world.db.expire_all()
    assert _vendor_ids(world.db) == sorted([world.a.id, world.b.id, world.c.id])
    assert _name(world, "b1") == "Acme GmbH & Co. KG"                        # nichts umbenannt
    assert world.db.get(Vendor, world.a.id).vat_id is None                   # nichts aufgefüllt
    assert world.db.query(VendorBankAccount).count() == 3


def test_vendor_list_has_document_counts(client, world):
    counts = {v["name"]: v["document_count"] for v in client.get("/api/vendors").json()}
    assert counts == {"ACME GmbH": 2, "Acme GmbH & Co. KG": 1, "Fremd AG": 1}


# ─── Namensvarianten: Belege müssen trotzdem mitwandern ─────────────────────────

def _variant_docs(w):
    def doc(name):
        d = crud.document.create(w.db, w.docs["a1"].batch_id, f"v{len(w.db.new)}{name[:3]}.pdf", 1)
        crud.document.save_extraction(w.db, d.id, {"supplier_name": name}, [], "{}")
        return d
    return {"case": doc("acme gmbh"), "space": doc("  ACME   GmbH "), "b_case": doc("ACME GMBH & CO. KG")}


def test_document_counts_ignore_case_and_whitespace(client, world):
    _variant_docs(world)
    by_name = {v["name"]: v["document_count"] for v in client.get("/api/vendors").json()}
    assert by_name["ACME GmbH"] == 4            # 2 exakt + 2 Schreibvarianten
    assert by_name["Acme GmbH & Co. KG"] == 2   # 1 exakt + 1 Variante


def test_merge_moves_documents_with_name_variants(client, world):
    variants = _variant_docs(world)
    r = client.post("/api/vendors/merge", json={"target_id": world.a.id, "source_ids": [world.b.id]})
    assert r.status_code == 200, r.text
    world.db.expire_all()
    for key in ("a1", "a2", "b1"):
        assert world.docs[key].extraction.vendor_id == "ACME GmbH"
    for d in variants.values():
        assert d.extraction.vendor_id == "ACME GmbH"
    assert world.docs["c1"].extraction.vendor_id == "Fremd AG"
    by_name = {v["name"]: v["document_count"] for v in client.get("/api/vendors").json()}
    assert by_name["ACME GmbH"] == 6


# ─── Reparatur verwaister Belege ─────────────────────────────────────────────────

def _orphan_doc(w, name, vat=None, iban=None):
    d = crud.document.create(w.db, w.docs["a1"].batch_id, f"o{len(w.db.new)}.pdf", 1)
    lief = {"name": name}
    if vat:
        lief["ust_id_nr"] = vat
    if iban:
        lief["bankverbindung"] = {"iban": iban}
    crud.document.save_extraction(w.db, d.id, {"supplier_name": name}, [], json.dumps({"lieferant": lief}))
    # save_extraction legt den Lieferanten an; verwaist wird der Beleg erst, wenn dieser (z. B. durch Zusammenführen) fehlt
    w.db.query(Vendor).filter(Vendor.name == name, Vendor.id.notin_([w.a.id, w.b.id, w.c.id])).delete()
    w.db.commit()
    return d


def test_orphans_are_found_with_suggestions(client, world):
    _orphan_doc(world, "Völlig anderer Name", vat="de 123456789")             # über USt-IdNr. → "Acme GmbH & Co. KG"
    _orphan_doc(world, "Weiterer Name", iban="DE89 3704 0044 0532 0130 00")  # über IBAN
    _orphan_doc(world, "Acme")                                                 # gleicher Name ohne Rechtsform → "ACME GmbH"
    _orphan_doc(world, "Unbekannt AG")                                      # kein Vorschlag
    r = client.get("/api/vendors/orphan-names").json()
    by = {o["name"]: o for o in r["orphans"]}
    assert r["total"] == 4
    assert by["Völlig anderer Name"]["suggestion"]["vendor_id"] == world.b.id
    assert by["Völlig anderer Name"]["suggestion"]["confidence"] == "high"
    assert by["Weiterer Name"]["suggestion"]["vendor_id"] == world.b.id
    assert by["Acme"]["suggestion"]["confidence"] in ("medium", "high")
    assert by["Unbekannt AG"]["suggestion"] is None
    assert "Acme GmbH" not in by                                         # Schreibvariante eines vorhandenen Lieferanten


def test_reassign_moves_documents_and_can_be_undone(client, world, sqlite_session_factory, monkeypatch):
    monkeypatch.setattr("app.services.change_log.SessionLocal", sqlite_session_factory)
    d1 = _orphan_doc(world, "Alter Name")
    d2 = _orphan_doc(world, "Alter Name")
    r = client.post("/api/vendors/reassign-names", json={"assignments": [{"name": "Alter Name", "target_id": world.a.id}]})
    assert r.status_code == 200 and r.json()["moved_documents"] == 2
    world.db.expire_all()
    assert d1.extraction.vendor_id == d2.extraction.vendor_id == "ACME GmbH"
    assert client.get("/api/vendors/orphan-names").json()["total"] == 0
    from app.models.data_change import DataChange
    assert world.db.query(DataChange).filter_by(run_id=r.json()["run_id"]).count() == 2
    from app.services import change_log
    change_log.start_undo(r.json()["run_id"], background=False)
    world.db.expire_all()
    assert d1.extraction.vendor_id == d2.extraction.vendor_id == "Alter Name"


def test_reassign_unknown_vendor_changes_nothing(client, world):
    d = _orphan_doc(world, "Alter Name")
    r = client.post("/api/vendors/reassign-names", json={"assignments": [{"name": "Alter Name", "target_id": 99999}]})
    assert r.status_code == 404
    world.db.expire_all()
    assert d.extraction.vendor_id == "Alter Name"


# ─── Excel-Export nach Lieferanten ─────────────────────────────────────────────

def _export(client, *ids):
    import io
    from openpyxl import load_workbook
    r = client.get("/api/vendors/export", params={"ids": list(ids)})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    return wb["Rechnungen"], wb["Positionen"], r


def test_vendor_export_contains_only_selected_vendors_documents(client, world):
    _variant_docs(world)                                    # 3 Belege mit Schreibvarianten (2× ACME, 1× Acme & Co.)
    ws, _, r = _export(client, world.a.id)
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=3)]            # Zeile 1 = Titel, 2 = Kopf
    header = [c.value for c in ws[2]]
    names = [row[header.index("Lieferant")] for row in rows]
    assert len(rows) == 4 and set(n.strip().casefold() for n in names) <= {"acme gmbh", "acme   gmbh"} or all("acme" in n.casefold() for n in names)
    assert not any("Fremd" in (n or "") for n in names)
    assert "xlsx" in r.headers["content-disposition"]
    assert ws["A1"].value == "Export: ACME GmbH"


def test_vendor_export_multiple_vendors_and_positions_sheet(client, world):
    ws, wsp, _ = _export(client, world.a.id, world.c.id)
    assert ws["A1"].value == "Export: 2 Lieferanten"
    docs_rows = [row for row in ws.iter_rows(min_row=3, values_only=True) if any(row)]
    assert len(docs_rows) == 3                              # a1, a2, c1


def test_vendor_export_errors(client, world):
    assert client.get("/api/vendors/export", params={"ids": [99999]}).status_code == 404
    empty = Vendor(name="Ohne Belege")
    world.db.add(empty)
    world.db.commit()
    r = client.get("/api/vendors/export", params={"ids": [empty.id]})
    assert r.status_code == 404 and "keine Belege" in r.json()["detail"]


def test_vendor_export_excludes_deleted_documents(client, world):
    world.docs["a2"].soft_deleted = True
    world.db.commit()
    ws, _, _ = _export(client, world.a.id)
    assert len([row for row in ws.iter_rows(min_row=3, values_only=True) if any(row)]) == 1


def test_vendor_export_skips_hidden_imports(client, world):
    from datetime import datetime, timezone
    world.docs["a2"].batch.hidden_at = datetime.now(timezone.utc)
    world.db.commit()
    r = client.get("/api/vendors/export", params={"ids": [world.a.id]})
    if world.docs["a1"].batch_id == world.docs["a2"].batch_id:      # alle Belege im ausgeblendeten Import
        assert r.status_code == 404
    else:
        assert r.status_code == 200
