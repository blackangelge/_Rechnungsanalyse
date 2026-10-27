"""Phase 5: schemagesteuerter Excel-Export — Spalten aus der Definition, echte Excel-Typen, Layouts, Mengenmodi, Filter, Vorlagen-API."""

import io
import json
from datetime import date, datetime, timezone
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import crud
from app.crud import schema_definition as sd
from app.database import get_db
from app.main import app
from app.models.export_template import ExportTemplate
from app.models.import_batch import ImportBatch
from app.services import schema_extraction as se
from app.services import schema_export, schema_models
from app.services.export_seed import seed_export_templates
from app.services.schema_seed import seed_initial_schema
from app.services.units_seed import seed_default_units
from tests.test_schema_extraction import SAMPLE, _Stub


@pytest.fixture()
def world(db):
    seed_default_units(db)
    version = seed_initial_schema(db)
    seed_export_templates(db)
    db.commit()
    schema_models.clear_cache()
    stub = _Stub(SAMPLE)
    cfg = SimpleNamespace(api_url=stub.url, api_key=None, model_name="m", max_tokens=100, temperature=0, reasoning="off")
    w = SimpleNamespace(db=db, version=version, stub=stub, cfg=cfg, docs=[])

    def add_doc(sample=None, company="T", year=2026, hidden=False):
        batch = ImportBatch(company_name=company, year=year, import_folder_path="/x", storage_folder_path="/s/x",
                            hidden_at=datetime.now(timezone.utc) if hidden else None)
        db.add(batch)
        db.flush()
        doc = crud.document.create(db, batch.id, f"re{len(w.docs)}.pdf", 1)
        doc.page_count = 2
        db.commit()
        stub.content = sample or SAMPLE
        out = se.run_sync(db, doc, [], cfg)
        assert out.save is not None, out.validation_errors or out.result.error
        w.docs.append(doc)
        return doc
    w.add_doc = add_doc
    yield w
    stub.close()


def _book(data: bytes):
    return load_workbook(io.BytesIO(data))


def _template(db, name_part):
    return next(t for t in db.query(ExportTemplate) if name_part in t.name)


def _export(w, template, **filters):
    return schema_export.build_workbook(w.db, sd.get_active(w.db), template, schema_export.Filters(**filters))


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_sheet_per_entity_columns_order_and_typed_cells(world):
    world.add_doc()
    wb = _book(_export(world, _template(world.db, "Blatt je Tabelle")))
    assert wb.sheetnames == ["Beleg", "Position", "Lieferant", "Umsatzsteuer"]
    head = [c.value for c in wb["Beleg"][1]]
    assert head[:5] == ["Datensatz-ID", "Beleg-ID", "Lieferant-ID", "Rechnungsnr.", "Rechnungsdatum"]       # Reihenfolge = Schema
    assert head[-2:] == ["Lieferant", "Lieferant USt-IdNr."]
    row = {h: c for h, c in zip(head, wb["Beleg"][2])}
    assert row["Rechnungsnr."].value == "RE-2024/0815"
    assert isinstance(row["Rechnungsdatum"].value, (date, datetime)) and row["Rechnungsdatum"].number_format == "DD.MM.YYYY"
    assert row["Rechnungsdatum"].value.year == 2025 and row["Rechnungsdatum"].value.month == 3
    assert row["Brutto (€)"].value == 4712.4 and row["Brutto (€)"].number_format == "#,##0.00"
    assert row["Skonto-Frist (Tage)"].value == 14 and row["Skonto-Frist (Tage)"].number_format == "0"
    assert row["Lieferant"].value == "ACME GmbH"
    assert row["Seiten"].value == 2 and row["Dateiname"].value == "re0.pdf"                                 # Systemfelder


def test_ids_connect_sheets_and_vendor_sheet_has_each_vendor_once(world):
    world.add_doc()
    world.add_doc()
    wb = _book(_export(world, _template(world.db, "Blatt je Tabelle")))
    beleg_ids = {r[0] for r in _rows(wb["Beleg"])[1:]}
    pos = _rows(wb["Position"])
    assert pos[0][:4] == ["Datensatz-ID", "Eltern-ID", "Beleg-ID", "Lieferant-ID"]
    assert len(pos) - 1 == 6 and {r[1] for r in pos[1:]} == beleg_ids                                       # jede Position zeigt auf einen Beleg
    vend = _rows(wb["Lieferant"])
    assert len(vend) == 2 and vend[1][1] == "ACME GmbH"                                                     # zwei Belege, ein Lieferant


def test_quantity_modes_original_vs_normalized(world):
    world.add_doc()
    t = _template(world.db, "Blatt je Tabelle")
    wb = _book(_export(world, t))
    head = [c.value for c in wb["Position"][1]]
    assert "Menge" in head and "Menge Einheit" in head
    kg_row = {h: c.value for h, c in zip(head, wb["Position"][3])}
    assert (kg_row["Menge"], kg_row["Menge Einheit"]) == (25880, "kg")                                      # wie gedruckt
    t.quantity_mode = "normalized"
    world.db.commit()
    wb = _book(_export(world, t))
    rows = {h: c.value for h, c in zip(head, wb["Position"][3])}
    assert (rows["Menge"], rows["Menge Einheit"]) == (25.88, "t")                                           # Basiseinheit
    unknown = {h: c.value for h, c in zip(head, wb["Position"][4])}
    assert unknown["Menge"] is None and unknown["Menge Einheit"] is None                                    # unbekannte Einheit: nicht erfunden
    t.quantity_mode = "original"
    world.db.commit()
    unknown = {h: c.value for h, c in zip(head, _book(_export(world, t))["Position"][4])}
    assert (unknown["Menge"], unknown["Menge Einheit"]) == (5, "Fass")                                      # Original bleibt sichtbar


def test_flat_layout_one_row_per_position_with_header_repeated(world):
    world.add_doc()
    world.add_doc()
    wb = _book(_export(world, _template(world.db, "Flach")))
    assert wb.sheetnames == ["Beleg + Position"]
    rows = _rows(wb.active)
    head = rows[0]
    assert len(rows) == 1 + 6 and head[:3] == ["Beleg-ID", "Datensatz-ID", "Lieferant-ID"]
    inv = head.index("Rechnungsnr.")
    assert {r[inv] for r in rows[1:]} == {"RE-2024/0815"}                                                   # Belegdaten wiederholt
    art = head.index("Artikelbezeichnung")
    assert [r[art] for r in rows[1:4]] == ["Schrauben M8", "Splitt", "Fassware"]
    menge = head.index("Menge")
    assert rows[2][menge] == 25.88 and rows[2][menge + 1] == "t"                                            # normiert (Vorlage)
    assert head.count("Währung") == 1 or any("Währung (" in h for h in head if h)                            # Kollision wird eindeutig


def test_flat_keeps_documents_without_positions(world):
    body = json.loads(json.dumps(SAMPLE))
    body["positionen"] = []
    world.add_doc(body)
    rows = _rows(_book(_export(world, _template(world.db, "Flach"))).active)
    assert len(rows) == 2 and rows[1][rows[0].index("Rechnungsnr.")] == "RE-2024/0815"
    assert rows[1][rows[0].index("Artikelbezeichnung")] is None


def test_field_selection_and_order_from_template(world):
    world.add_doc()
    t = ExportTemplate(name="Schmal", layout="sheet_per_entity", quantity_mode="original",
                       config={"entities": [{"entity": "beleg", "fields": ["gesamtbetrag_brutto", "rechnungsnummer"]}]})
    world.db.add(t)
    world.db.commit()
    wb = _book(_export(world, t))
    assert wb.sheetnames == ["Beleg"]                                       # nur gewählte Entität
    assert [c.value for c in wb["Beleg"][1]] == ["Datensatz-ID", "Beleg-ID", "Lieferant-ID", "Brutto (€)", "Rechnungsnr."]


def test_filters_batches_documents_vendors_and_hidden(world):
    a = world.add_doc(company="A")
    b = world.add_doc(company="B", hidden=True)
    other = json.loads(json.dumps(SAMPLE))
    other["lieferant"].update({"name": "Andere AG", "iban": None, "ust_id_nr": "DE000"})
    c = world.add_doc(other, company="C")
    t = _template(world.db, "Blatt je Tabelle")

    def ids(**f):
        return sorted(r[1] for r in _rows(_book(_export(world, t, **f))["Beleg"])[1:])
    assert ids() == sorted([a.id, c.id])                                   # ausgeblendete Imports standardmäßig draußen
    assert ids(include_hidden=True) == sorted([a.id, b.id, c.id])
    assert ids(batch_ids=[b.batch_id]) == [b.id]                           # ausdrücklich gewählter Import gewinnt
    assert ids(document_ids=[c.id]) == [c.id]
    from app.models.schema_def import ExtractedRecord
    vid = world.db.query(ExtractedRecord).filter_by(document_id=c.id, parent_record_id=None).one().vendor_id
    assert ids(vendor_ids=[vid]) == [c.id]
    with pytest.raises(schema_export.ExportError, match="keine Datensätze"):
        _export(world, t, document_ids=[999999])


def test_soft_deleted_documents_are_excluded(world):
    a = world.add_doc()
    b = world.add_doc()
    b.soft_deleted = True
    world.db.commit()
    ids = [r[1] for r in _rows(_book(_export(world, _template(world.db, "Blatt je Tabelle")))["Beleg"])[1:]]
    assert ids == [a.id]


def test_old_version_records_stay_exportable(world):
    old = world.add_doc()
    draft = sd.create_draft_from(world.db, world.version)
    beleg = next(e for e in draft.entities if e.name == "beleg")
    sd.add_field(world.db, beleg, "kostenstelle", "Kostenstelle", "text")
    sd.activate(world.db, draft)
    world.db.commit()
    schema_models.clear_cache()
    body = json.loads(json.dumps(SAMPLE))
    body["kostenstelle"] = "KS-7"
    new = world.add_doc(body)
    wb = _book(_export(world, _template(world.db, "Blatt je Tabelle")))
    head = [c.value for c in wb["Beleg"][1]]
    ks = head.index("Kostenstelle")
    by_doc = {r[1]: r[ks] for r in _rows(wb["Beleg"])[1:]}
    assert by_doc == {old.id: None, new.id: "KS-7"}                        # Altdaten bleiben lesbar, neue Spalte leer


def test_text_is_sanitized_and_excel_types_survive_odd_values(world):
    body = json.loads(json.dumps(SAMPLE))
    body["zahlungsbedingungen"] = "Zahlbar\x00 sofort\x07"
    body["positionen"][0]["artikelbezeichnung"] = "=SUMME(A1:A2)"
    world.add_doc(body)
    wb = _book(_export(world, _template(world.db, "Blatt je Tabelle")))
    head = [c.value for c in wb["Beleg"][1]]
    assert wb["Beleg"][2][head.index("Zahlungsbedingungen")].value == "Zahlbar sofort"
    assert wb["Position"][2][[c.value for c in wb["Position"][1]].index("Artikelbezeichnung")].value == "=SUMME(A1:A2)"


# ─── API ─────────────────────────────────────────────────────────────────────

@pytest.fixture()
def client(sqlite_session_factory):
    def _db():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _db
    yield TestClient(app)
    app.dependency_overrides.pop(get_db, None)


def test_template_api_crud_and_validation(client, world):
    listing = client.get("/api/export-templates").json()
    assert len(listing) == 2 and listing[0]["layout"] in ("flat", "sheet_per_entity")
    ok = {"name": "Neu", "layout": "sheet_per_entity", "quantity_mode": "normalized",
          "config": {"entities": [{"entity": "beleg", "fields": ["rechnungsnummer"]}]}}
    r = client.post("/api/export-templates", json=ok)
    assert r.status_code == 201
    tid = r.json()["id"]
    assert client.post("/api/export-templates", json=ok).status_code == 409                       # Name doppelt
    for bad, msg in (
        ({**ok, "name": "A", "layout": "kreuz"}, "Layout"),
        ({**ok, "name": "B", "quantity_mode": "x"}, "Mengenmodus"),
        ({**ok, "name": "C", "config": {"entities": [{"entity": "gibtsnicht"}]}}, "gibt es im aktiven Schema nicht"),
        ({**ok, "name": "D", "config": {"entities": [{"entity": "beleg", "fields": ["nope"]}]}}, "Feld"),
        ({**ok, "name": "E", "config": {"vendor": ["iban", "blz"]}}, "Lieferantenfeld"),
        ({**ok, "name": "F", "layout": "flat", "config": {"flat": {"root": "lieferant", "child": "position"}}}, "Eltern-Kind"),
    ):
        r = client.post("/api/export-templates", json=bad)
        assert r.status_code == 422 and msg in r.json()["detail"], (bad, r.text)
    assert client.put(f"/api/export-templates/{tid}", json={**ok, "name": "Umbenannt"}).json()["name"] == "Umbenannt"
    cols = client.get(f"/api/export-templates/{tid}/columns").json()
    assert cols == [{"entity": "beleg", "display_name": "Beleg", "columns": [{"header": "Rechnungsnr.", "kind": "text", "format": None}]}]
    assert client.delete(f"/api/export-templates/{tid}").status_code == 204
    assert client.get(f"/api/export-templates/{tid}").status_code == 404


def test_export_endpoint_streams_xlsx_or_explains(client, world):
    world.add_doc()
    tid = _template(world.db, "Blatt je Tabelle").id
    r = client.get(f"/api/export-templates/{tid}/export")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and "attachment" in r.headers["content-disposition"]
    assert _book(r.content).sheetnames == ["Beleg", "Position", "Lieferant", "Umsatzsteuer"]
    r = client.get(f"/api/export-templates/{tid}/export", params={"document_ids": [987654]})
    assert r.status_code == 404 and "keine Datensätze" in r.json()["detail"]
    assert client.get("/api/export-templates/999/export").status_code == 404


def test_seed_is_idempotent(world):
    assert seed_export_templates(world.db) is False
