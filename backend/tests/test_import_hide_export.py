"""Imports ausblenden/einblenden und gemeinsamer Excel-Export."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import crud
from app.database import get_db
from app.main import app
from app.models.import_batch import ImportBatch


@pytest.fixture()
def client(sqlite_session_factory):
    def _get_db():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _get_db
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture()
def world(db):
    """Drei Imports (A 2024, B 2025, C 2025) mit je 2 Belegen (B1 und C2 mit Positionen)."""
    class W:
        pass
    w = W()
    w.db, w.batches, w.docs = db, {}, {}
    for name, year in (("A", 2024), ("B", 2025), ("C", 2025)):
        b = ImportBatch(company_name=f"Firma {name}", year=year, import_folder_path="/x", storage_folder_path=f"/s/{name}", status="done")
        db.add(b)
        db.commit()
        w.batches[name] = b
        for i in (1, 2):
            d = crud.document.create(db, b.id, f"{name}{i}.pdf", 1)
            positions = [{"product_description": f"Artikel {name}{i}", "quantity": 2, "unit": "kg",
                          "unit_price": 1.5, "total_price": 3.0}] if (name, i) in (("B", 1), ("C", 2)) else []
            crud.document.save_extraction(
                db, d.id, {"supplier_name": f"Lief {name}", "total_amount": 10 * i, "invoice_number": f"R-{name}{i}"},
                positions, "{}",
            )
            w.docs[f"{name}{i}"] = d
    return w


def _hide(client, w, *names, hidden=True):
    return client.post("/api/imports/visibility", json={"ids": [w.batches[n].id for n in names], "hidden": hidden})


# ─── Ausblenden ───────────────────────────────────────────────────────────────

def test_hide_and_unhide_roundtrip(client, world):
    ids = lambda **p: sorted(b["id"] for b in client.get("/api/imports", params=p).json())     # noqa: E731
    all_ids = sorted(b.id for b in world.batches.values())
    assert ids() == all_ids
    r = _hide(client, world, "A", "B")
    assert r.status_code == 200 and r.json()["updated"] == 2
    assert ids() == [world.batches["C"].id]                                    # standardmäßig ausgeblendet
    listing = client.get("/api/imports", params={"include_hidden": True}).json()
    assert sorted(b["id"] for b in listing) == all_ids
    hidden = {b["id"]: b["hidden_at"] for b in listing}
    assert hidden[world.batches["A"].id] and hidden[world.batches["C"].id] is None
    r = _hide(client, world, "A", hidden=False)
    assert r.json()["updated"] == 1 and ids() == sorted([world.batches["A"].id, world.batches["C"].id])


def test_hide_is_idempotent_and_keeps_timestamp(client, world):
    _hide(client, world, "A")
    first = {b["id"]: b["hidden_at"] for b in client.get("/api/imports", params={"include_hidden": True}).json()}
    again = _hide(client, world, "A")
    assert again.json()["updated"] == 0                                         # nichts geändert
    second = {b["id"]: b["hidden_at"] for b in client.get("/api/imports", params={"include_hidden": True}).json()}
    assert first == second
    assert _hide(client, world, "B", hidden=False).json()["updated"] == 0       # war sichtbar


def test_hide_validation_and_unknown_ids(client, world):
    assert client.post("/api/imports/visibility", json={"ids": [], "hidden": True}).status_code == 422
    r = client.post("/api/imports/visibility", json={"ids": [world.batches["A"].id, 9999], "hidden": True})
    assert r.status_code == 404 and "#9999" in r.json()["detail"]
    assert client.get("/api/imports").json().__len__() == 3                     # nichts wurde ausgeblendet


def test_hiding_does_not_delete_anything(client, world):
    _hide(client, world, "A", "B", "C")
    world.db.expire_all()
    assert world.db.query(ImportBatch).count() == 3
    assert client.get(f"/api/imports/{world.batches['A'].id}").status_code == 200       # Detailseite weiter erreichbar
    assert len(client.get("/api/documents", params={"include_hidden": True}).json()) == 6


# ─── Belege ausgeblendeter Imports ────────────────────────────────────────────

def test_documents_of_hidden_imports_are_excluded_by_default(client, world):
    _hide(client, world, "A")
    names = lambda **p: sorted(d["original_filename"] for d in client.get("/api/documents", params=p).json())     # noqa: E731
    assert names() == ["B1.pdf", "B2.pdf", "C1.pdf", "C2.pdf"]
    assert len(names(include_hidden=True)) == 6
    # Zähler, ID-Liste und Paginierung sehen dasselbe
    r = client.get("/api/documents", params={"limit": 2})
    assert r.headers["x-total-count"] == "4"
    assert len(client.get("/api/documents/ids").json()) == 4
    assert len(client.get("/api/documents/ids", params={"include_hidden": True}).json()) == 6


def test_explicit_import_or_doc_id_bypasses_hiding(client, world):
    _hide(client, world, "A", "B")
    a = world.batches["A"].id
    got = client.get("/api/documents", params={"batch_ids": a}).json()
    assert sorted(d["original_filename"] for d in got) == ["A1.pdf", "A2.pdf"]          # bewusst gewählter Import
    one = client.get("/api/documents", params={"doc_id": world.docs["A1"].id}).json()
    assert [d["original_filename"] for d in one] == ["A1.pdf"]                          # Link aus den Logs
    # ohne Auswahl bleiben A und B draußen
    assert sorted(d["original_filename"] for d in client.get("/api/documents").json()) == ["C1.pdf", "C2.pdf"]


def test_filter_options_follow_visibility(client, world):
    _hide(client, world, "A")
    o = client.get("/api/documents/filter-options").json()
    assert o["years"] == [2025] and o["companies"] == ["Firma B", "Firma C"]
    o = client.get("/api/documents/filter-options", params={"include_hidden": True}).json()
    assert o["years"] == [2025, 2024]


# ─── Gemeinsamer Excel-Export ─────────────────────────────────────────────────

def _wb(resp):
    return load_workbook(io.BytesIO(resp.content))


def test_multi_export_has_import_column_and_all_rows(client, world):
    ids = [world.batches["B"].id, world.batches["C"].id]
    r = client.get("/api/imports/export", params={"ids": ids})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Export_2_Imports_" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.xlsx"')
    wb = _wb(r)
    assert wb.sheetnames == ["Rechnungen", "Positionen"]
    inv, pos = wb["Rechnungen"], wb["Positionen"]
    assert inv["A1"].value == "Export: 2 Imports"
    headers = [c.value for c in inv[2]]
    assert headers[-1] == "Import" and "Rechnungsnr." in headers
    ri = headers.index("Rechnungsnr.")
    rows = [[c.value for c in r_] for r_ in inv.iter_rows(min_row=3)]
    assert len(rows) == 4
    assert {(r_[ri], r_[-1]) for r_ in rows} == {
        ("R-B1", f"Firma B 2025 (#{ids[0]})"), ("R-B2", f"Firma B 2025 (#{ids[0]})"),
        ("R-C1", f"Firma C 2025 (#{ids[1]})"), ("R-C2", f"Firma C 2025 (#{ids[1]})"),
    }
    pheaders = [c.value for c in pos[1]]
    assert pheaders[-1] == "Import"
    ai = pheaders.index("Artikelbezeichnung")
    prows = [[c.value for c in r_] for r_ in pos.iter_rows(min_row=2)]
    assert sorted((r_[ai], r_[-1]) for r_ in prows) == [
        ("Artikel B1", f"Firma B 2025 (#{ids[0]})"), ("Artikel C2", f"Firma C 2025 (#{ids[1]})")]


def test_multi_export_excludes_soft_deleted_documents(client, world):
    crud.document.soft_delete(world.db, world.docs["B1"].id)
    ids = [world.batches["B"].id, world.batches["C"].id]
    wb = _wb(client.get("/api/imports/export", params={"ids": ids}))
    assert wb["Rechnungen"].max_row - 2 == 3


def test_export_with_single_id_keeps_the_classic_layout(client, world):
    b = world.batches["B"]
    r = client.get("/api/imports/export", params={"ids": [b.id]})
    wb = _wb(r)
    assert "Export_Firma_B_2025.xlsx" in r.headers["content-disposition"]
    assert "Import" not in [c.value for c in wb["Rechnungen"][2]]                  # keine Import-Spalte
    classic = _wb(client.get(f"/api/imports/{b.id}/export"))
    assert [c.value for c in classic["Rechnungen"][2]] == [c.value for c in wb["Rechnungen"][2]]


def test_export_works_for_hidden_imports_and_rejects_unknown(client, world):
    _hide(client, world, "A")
    ids = [world.batches["A"].id, world.batches["B"].id]
    assert client.get("/api/imports/export", params={"ids": ids}).status_code == 200
    r = client.get("/api/imports/export", params={"ids": [world.batches["A"].id, 4711]})
    assert r.status_code == 404 and "#4711" in r.json()["detail"]
    assert client.get("/api/imports/export").status_code == 422                    # ids fehlt
    dup = client.get("/api/imports/export", params={"ids": [ids[0], ids[0]]})
    assert dup.status_code == 200 and "Export_Firma_A_2024" in dup.headers["content-disposition"]
