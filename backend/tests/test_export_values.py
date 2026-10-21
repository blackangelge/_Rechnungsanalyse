"""Excel-Export gibt die gespeicherten Datenbankwerte aus — ohne Berechnung (kein Menge × Preis) und ohne Umrechnung."""

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import crud
from app.database import get_db
from app.main import app
from app.models.import_batch import ImportBatch


@pytest.fixture()
def client(sqlite_session_factory):
    def _get_db():
        s = sqlite_session_factory()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _get_db
    yield TestClient(app)
    app.dependency_overrides.clear()


def _sheet(client, batch_id):
    r = client.get(f"/api/imports/{batch_id}/export")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Positionen"]
    header = [c.value for c in ws[1]]
    return header, [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def test_positions_sheet_contains_stored_values_only(client, db):
    b = ImportBatch(company_name="F", year=2025, import_folder_path="/x", storage_folder_path="/s", status="done")
    db.add(b)
    db.commit()
    d = crud.document.create(db, b.id, "a.pdf", 1)
    crud.document.save_extraction(
        db, d.id, {"supplier_name": "L", "invoice_number": "R-1"},
        [{"product_description": "KAS", "quantity": 25500, "unit": "kg", "unit_price": 412.0, "total_price": 10506.0, "discount": None},
         {"product_description": "Teil", "quantity": 3, "unit": "Stück", "unit_price": 0.4125, "total_price": 1.24}],
        json.dumps({"positionen": [{}, {}]}),
    )
    header, rows = _sheet(client, b.id)
    col = {h: i for i, h in enumerate(header)}
    assert "Gesamtpreis (€)" in header and "Einzelpreis (€)" in header and "EP Brutto (€)" not in header
    r0, r1 = rows[0], rows[1]
    assert (r0[col["Menge"]], r0[col["Einheit"]], r0[col["Einzelpreis (€)"]], r0[col["Gesamtpreis (€)"]]) == (25500, "kg", 412.0, 10506.0)
    # Gesamtpreis ist der gespeicherte Wert — NICHT Menge × Einzelpreis (25.500 × 412 = 10.506.000)
    assert r0[col["Gesamtpreis (€)"]] != r0[col["Menge"]] * r0[col["Einzelpreis (€)"]]
    assert r1[col["Einzelpreis (€)"]] == 0.4125 and r1[col["Gesamtpreis (€)"]] == 1.24                # keine Rundung, kein Nachrechnen
    ws = load_workbook(io.BytesIO(client.get(f"/api/imports/{b.id}/export").content))["Positionen"]
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)      # keine Excel-Formeln
