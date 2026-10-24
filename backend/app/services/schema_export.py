"""
Schemagesteuerter Excel-Export: Spalten, Reihenfolge, Überschriften und Zellformate kommen aus der Schema-Definition,
Auswahl und Layout aus einer Export-Vorlage (`models/export_template.py`).

Layouts
  sheet_per_entity — ein Blatt je gewählter Entität; verbunden über ID-Spalten (Datensatz-ID, Eltern-ID, Beleg-ID, Lieferant-ID).
                     Stammdaten-Entitäten (Lieferant) erscheinen als Blatt mit den im Export vorkommenden Lieferanten.
  flat             — eine Zeile je Kindsatz (Position) mit wiederholten Daten des Elternsatzes (Beleg); Belege ohne Kinder bekommen eine Zeile.

Echte Excel-Typen: Datum als Datum, Zahlen als Zahl (Format aus `output_format` des Feldes), Wahrheitswerte als Boolean.
Mengenfelder: je Feld zwei Spalten (Wert + Einheit) — wahlweise Originalwert/-einheit (wie gedruckt) oder normiert (Basiseinheit).
Schreibt im openpyxl-`write_only`-Modus über Blöcke von 300 Belegen (Speicherbedarf unabhängig von der Datenmenge).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Iterator

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.models.document import Document
from app.models.export_template import LAYOUTS, QUANTITY_MODES, ExportTemplate
from app.models.import_batch import ImportBatch
from app.models.schema_def import ExtractedRecord, SchemaEntity, SchemaField, SchemaVersion
from app.models.vendor import Vendor

CHUNK = 300
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
VENDOR_COLUMNS = {                      # Schlüssel (wie die Felder der Lieferant-Entität) → Vendor-Attribut
    "name": "name", "strasse": "street", "plz": "postal_code", "ort": "city", "land": "country",
    "hrb_nummer": "hrb_number", "steuernummer": "tax_number", "ust_id_nr": "vat_id",
}
_DEFAULT_FORMATS = {"decimal": "#,##0.00", "integer": "0", "date": "DD.MM.YYYY"}


class ExportError(ValueError):
    """Vorlage oder Auswahl nicht nutzbar."""


# ─── Vorlage prüfen ──────────────────────────────────────────────────────────

def validate_template(version: SchemaVersion, layout: str, quantity_mode: str, config: dict) -> None:
    if layout not in LAYOUTS:
        raise ExportError(f"Unbekanntes Layout: {layout}")
    if quantity_mode not in QUANTITY_MODES:
        raise ExportError(f"Unbekannter Mengenmodus: {quantity_mode}")
    by_name = {e.name: e for e in version.entities}
    for sel in config.get("entities", []):
        ent = by_name.get(sel.get("entity"))
        if ent is None:
            raise ExportError(f"Entität „{sel.get('entity')}“ gibt es im aktiven Schema nicht")
        names = {f.name for f in ent.fields}
        for fname in sel.get("fields") or []:
            if fname not in names:
                raise ExportError(f"Feld „{fname}“ gibt es in {ent.name} nicht")
    for key in config.get("vendor", []):
        if key not in VENDOR_COLUMNS and key != "iban":
            raise ExportError(f"Lieferantenfeld „{key}“ ist nicht verfügbar")
    if layout == "flat":
        root, child = flat_entities(version, config)
        if not any(r.kind == "parent_child" and r.source_entity_id == root.id and r.target_entity_id == child.id for r in version.relations):
            raise ExportError(f"Zwischen {root.name} und {child.name} gibt es keine Eltern-Kind-Beziehung")


def flat_entities(version: SchemaVersion, config: dict) -> tuple[SchemaEntity, SchemaEntity]:
    by_name = {e.name: e for e in version.entities}
    spec = config.get("flat") or {}
    if spec.get("root") and spec.get("child"):
        try:
            return by_name[spec["root"]], by_name[spec["child"]]
        except KeyError as exc:
            raise ExportError(f"Entität {exc} gibt es im aktiven Schema nicht") from exc
    rel = next((r for r in version.relations if r.kind == "parent_child"), None)
    if rel is None:
        raise ExportError("Das Schema hat keine Eltern-Kind-Beziehung für das flache Layout")
    by_id = {e.id: e for e in version.entities}
    return by_id[rel.source_entity_id], by_id[rel.target_entity_id]


# ─── Spalten ─────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class Col:
    key: str
    header: str
    kind: str          # text | integer | decimal | date | boolean | qty_value | qty_unit | id
    fmt: str | None = None


def entity_columns(entity: SchemaEntity, selected: list[str] | None, mode: str) -> list[Col]:
    by_name = {f.name: f for f in entity.fields}
    if selected:
        fields = [by_name[n] for n in selected if n in by_name]
    else:
        fields = [f for f in entity.fields if not f.is_locked or f.system_function]
    cols: list[Col] = []
    for f in fields:
        if f.data_type == "quantity":
            cols.append(Col(f.name, f.display_name, "qty_value", f.output_format or ("#,##0.######" if mode == "normalized" else "#,##0.####")))
            cols.append(Col(f.name, f"{f.display_name} Einheit", "qty_unit"))
        else:
            cols.append(Col(f.name, f.display_name, f.data_type, f.output_format or _DEFAULT_FORMATS.get(f.data_type)))
    return cols


_RULE_STATUS_LABELS = {"ok": "in Ordnung", "warn": "Hinweis", "needs_review": "Prüfen!", "unchecked": "nicht prüfbar"}


def rule_text(data: dict) -> str | None:
    """Prüfregel-Ergebnis eines Datensatzes als Text: „OK“, „weicht ab: <Regel>; …“ (+ „nicht prüfbar“), leer = nicht geprüft."""
    checks = data.get("_checks") or []
    if not checks:
        return None
    failed = [c.get("name") for c in checks if c.get("status") == "failed"]
    skipped = sum(1 for c in checks if c.get("status") == "skipped")
    parts = ["weicht ab: " + "; ".join(str(n) for n in failed)] if failed else ["OK"]
    if skipped:
        parts.append(f"{skipped} nicht prüfbar")
    return _text(", ".join(parts))


def rule_status_text(data: dict) -> str | None:
    v = data.get("_validation")
    return _RULE_STATUS_LABELS.get(v.get("status")) if isinstance(v, dict) else None


def _vendor_cols(keys: list[str]) -> list[Col]:
    labels = {"name": "Lieferant", "strasse": "Lieferant Straße", "plz": "Lieferant PLZ", "ort": "Lieferant Ort", "land": "Lieferant Land",
              "hrb_nummer": "Lieferant HRB-Nr.", "steuernummer": "Lieferant Steuernr.", "ust_id_nr": "Lieferant USt-IdNr.", "iban": "Lieferant IBAN"}
    return [Col(f"vendor:{k}", labels[k], "text") for k in keys]


# ─── Werte ───────────────────────────────────────────────────────────────────

def _text(v) -> str | None:
    if v is None:
        return None
    s = _CTRL.sub("", str(v))
    return s[:32000] if s else None


def cell_value(col: Col, data: dict, mode: str):
    v = data.get(col.key)
    if col.kind in ("qty_value", "qty_unit"):
        if not isinstance(v, dict):
            return None
        if col.kind == "qty_value":
            return v.get("original_value" if mode == "original" else "value")
        return _text(v.get("original_unit" if mode == "original" else "unit"))
    if v is None:
        return None
    if col.kind == "date":
        try:
            return date.fromisoformat(str(v)[:10])
        except ValueError:
            return _text(v)
    if col.kind in ("integer", "decimal"):
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else _text(v)
    if col.kind == "boolean":
        return bool(v)
    return _text(v)


# ─── Daten ───────────────────────────────────────────────────────────────────

@dataclass
class Filters:
    batch_ids: list[int] | None = None
    document_ids: list[int] | None = None
    vendor_ids: list[int] | None = None
    include_hidden: bool = False


def _document_ids_query(db: Session, f: Filters):
    q = (
        db.query(Document.id).join(ExtractedRecord, ExtractedRecord.document_id == Document.id)
        .filter(ExtractedRecord.parent_record_id.is_(None), Document.soft_deleted.is_(False))
    )
    if f.batch_ids:
        q = q.filter(Document.batch_id.in_(f.batch_ids))
    elif not f.include_hidden:
        q = q.join(ImportBatch, ImportBatch.id == Document.batch_id).filter(ImportBatch.hidden_at.is_(None))
    if f.document_ids:
        q = q.filter(Document.id.in_(f.document_ids))
    if f.vendor_ids:
        q = q.filter(ExtractedRecord.vendor_id.in_(f.vendor_ids))
    return q


def iter_chunks(db: Session, f: Filters) -> Iterator[list[ExtractedRecord]]:
    """Alle Datensätze von je CHUNK Belegen (Keyset über die Beleg-ID)."""
    last = 0
    base = _document_ids_query(db, f)
    while True:
        ids = [r[0] for r in base.filter(Document.id > last).order_by(Document.id).distinct().limit(CHUNK)]
        if not ids:
            return
        last = ids[-1]
        recs = db.query(ExtractedRecord).filter(ExtractedRecord.document_id.in_(ids)).order_by(
            ExtractedRecord.document_id, ExtractedRecord.parent_record_id.isnot(None), ExtractedRecord.position_index, ExtractedRecord.id).all()
        yield recs
        for r in recs:                       # nur die eigenen Objekte freigeben — die Sitzung gehört dem Aufrufer
            db.expunge(r)


def count_documents(db: Session, f: Filters) -> int:
    return _document_ids_query(db, f).distinct().count()


# ─── Schreiben ───────────────────────────────────────────────────────────────

_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Blatt"
    t, n = base, 2
    while t in used:
        t = f"{base[:28]} {n}"
        n += 1
    used.add(t)
    return t


def _write_sheet_header(ws, headers: list[str]) -> None:
    cells = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.fill = _HEAD_FONT, _HEAD_FILL
        cells.append(c)
    ws.append(cells)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[_letter(i)].width = max(10, min(40, len(h) + 4))


def _letter(i: int) -> str:
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _row(ws, values: list, cols: list[Col | None]):
    out = []
    for v, col in zip(values, cols):
        if v is None or col is None or not col.fmt or isinstance(v, str):
            out.append(v)
            continue
        c = WriteOnlyCell(ws, value=v)
        c.number_format = col.fmt
        out.append(c)
    return out


def _vendor_values(vendor: Vendor | None, keys: list[str]) -> list:
    if vendor is None:
        return [None] * len(keys)
    out = []
    for k in keys:
        if k == "iban":
            out.append(_text(next((b.iban for b in vendor.bank_accounts if b.iban), None)))
        else:
            out.append(_text(getattr(vendor, VENDOR_COLUMNS[k])))
    return out


def build_workbook(db: Session, version: SchemaVersion, template: ExportTemplate, filters: Filters) -> bytes:
    validate_template(version, template.layout, template.quantity_mode, template.config)
    if count_documents(db, filters) == 0:                 # vor dem Anlegen der Arbeitsmappe (sonst bleibt ein halber Writer liegen)
        raise ExportError("Zu dieser Auswahl gibt es keine Datensätze.")
    mode = template.quantity_mode
    cfg = template.config or {}
    by_name = {e.name: e for e in version.entities}
    entity_name_by_id = {e.id: e.name for e in db.query(SchemaEntity).all()}      # auch Entitäten älterer Versionen
    selection = cfg.get("entities") or [{"entity": e.name} for e in version.entities]
    vendor_keys = list(cfg.get("vendor") or [])
    rule_check = bool(cfg.get("rule_check"))             # Spalten mit dem Ergebnis der Prüfregeln

    wb = Workbook(write_only=True)
    used: set[str] = set()
    vendors_seen: set[int] = set()
    n_docs = 0

    if template.layout == "sheet_per_entity":
        sheets: dict[str, dict] = {}
        master_sheets: dict[str, dict] = {}
        for sel in selection:
            ent = by_name[sel["entity"]]
            cols = entity_columns(ent, sel.get("fields"), mode)
            if ent.is_master:
                ws = wb.create_sheet(_sheet_title(ent.display_name, used))
                key_cols = [Col(f"vendor:{c.key}", c.header, "text") for c in cols if c.key in VENDOR_COLUMNS or c.key == "iban"]
                master_sheets[ent.name] = {"ws": ws, "cols": key_cols}
                _write_sheet_header(ws, ["Lieferant-ID"] + [c.header for c in key_cols])
                continue
            is_child = any(r.kind == "parent_child" and r.target_entity_id == ent.id for r in version.relations)
            extra = _vendor_cols(vendor_keys) if not is_child and ent.name == _root_name(version) else []
            ids = ["Datensatz-ID"] + (["Eltern-ID"] if is_child else []) + ["Beleg-ID", "Lieferant-ID"]
            ws = wb.create_sheet(_sheet_title(ent.display_name, used))
            is_root = not is_child and ent.name == _root_name(version)
            rule_headers = (["Prüfregel-Status", "Prüfregeln"] if is_root else ["Prüfregeln"]) if rule_check else []
            _write_sheet_header(ws, ids + [c.header for c in cols] + [c.header for c in extra] + rule_headers)
            sheets[ent.name] = {"ws": ws, "cols": cols, "extra": extra, "child": is_child, "rules": rule_headers, "root": is_root}
        vendor_cache: dict[int, Vendor] = {}
        for recs in iter_chunks(db, filters):
            n_docs += len({r.document_id for r in recs})
            if vendor_keys:
                need = {r.vendor_id for r in recs if r.vendor_id and r.vendor_id not in vendor_cache}
                if need:
                    vendor_cache.update({v.id: v for v in db.query(Vendor).filter(Vendor.id.in_(need)).all()})
            for r in recs:
                if r.vendor_id:
                    vendors_seen.add(r.vendor_id)
                s = sheets.get(entity_name_by_id.get(r.entity_id, ""))
                if s is None:
                    continue
                ids = [r.id] + ([r.parent_record_id] if s["child"] else []) + [r.document_id, r.vendor_id]
                vals = [cell_value(c, r.data or {}, mode) for c in s["cols"]]
                extra_vals = _vendor_values(vendor_cache.get(r.vendor_id), vendor_keys) if s["extra"] else []
                if s["rules"]:
                    extra_vals = extra_vals + ([rule_status_text(r.data or {})] if s["root"] else []) + [rule_text(r.data or {})]
                s["ws"].append(_row(s["ws"], ids + vals + extra_vals, [None] * len(ids) + s["cols"] + [None] * len(extra_vals)))
        for m in master_sheets.values():
            keys = [c.key.split(":", 1)[1] for c in m["cols"]]
            for v in db.query(Vendor).filter(Vendor.id.in_(vendors_seen)).order_by(Vendor.id) if vendors_seen else []:
                m["ws"].append([v.id] + _vendor_values(v, keys))
    else:
        root, child = flat_entities(version, cfg)
        sel_by = {s["entity"]: s for s in selection}
        root_cols = entity_columns(root, (sel_by.get(root.name) or {}).get("fields"), mode)
        child_cols = entity_columns(child, (sel_by.get(child.name) or {}).get("fields"), mode)
        v_cols = _vendor_cols(vendor_keys)
        root_rules = ["Prüfregel-Status", "Prüfregeln (Beleg)"] if rule_check else []
        headers = ["Beleg-ID", "Datensatz-ID", "Lieferant-ID"] + [c.header for c in root_cols] + [c.header for c in v_cols] + root_rules + [f"{child.display_name}-ID"]
        seen = set(headers)
        child_headers = []
        for c in child_cols:
            h = c.header if c.header not in seen else f"{c.header} ({child.display_name})"
            seen.add(h)
            child_headers.append(h)
        ws = wb.create_sheet(_sheet_title(root.display_name + " + " + child.display_name, used))
        _write_sheet_header(ws, headers + child_headers + ([f"Prüfregeln ({child.display_name})"] if rule_check else []))
        vendor_cache = {}
        fmts = [None] * 3 + root_cols + [None] * (len(v_cols) + len(root_rules)) + [None] + child_cols
        for recs in iter_chunks(db, filters):
            n_docs += len({r.document_id for r in recs})
            if vendor_keys:
                need = {r.vendor_id for r in recs if r.vendor_id and r.vendor_id not in vendor_cache}
                if need:
                    vendor_cache.update({v.id: v for v in db.query(Vendor).filter(Vendor.id.in_(need)).all()})
            kids: dict[int, list[ExtractedRecord]] = {}
            for r in recs:
                if r.parent_record_id is not None and entity_name_by_id.get(r.entity_id) == child.name:
                    kids.setdefault(r.parent_record_id, []).append(r)
            for r in recs:
                if r.parent_record_id is not None or entity_name_by_id.get(r.entity_id) != root.name:
                    continue
                head = [r.document_id, r.id, r.vendor_id] + [cell_value(c, r.data or {}, mode) for c in root_cols] \
                    + _vendor_values(vendor_cache.get(r.vendor_id), vendor_keys) \
                    + ([rule_status_text(r.data or {}), rule_text(r.data or {})] if rule_check else [])
                for k in kids.get(r.id) or [None]:
                    tail = [k.id] + [cell_value(c, k.data or {}, mode) for c in child_cols] if k is not None else [None] * (1 + len(child_cols))
                    if rule_check:
                        tail = tail + [rule_text(k.data or {}) if k is not None else None]
                    ws.append(_row(ws, head + tail, fmts))

    if n_docs == 0:
        raise ExportError("Zu dieser Auswahl gibt es keine Datensätze.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _root_name(version: SchemaVersion) -> str:
    from app.services.schema_models import root_entity
    return root_entity(version).name
